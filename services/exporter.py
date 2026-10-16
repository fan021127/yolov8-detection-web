import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


def export_records_to_excel(records):
    """
    将检测记录导出为 Excel 文件（内存中生成，返回 bytes）
    records: DetectionRecord.list_all_for_export 的返回值
    """
    wb = Workbook()

    # ---------- Sheet1: 检测记录汇总 ----------
    ws1 = wb.active
    ws1.title = '检测记录汇总'

    headers = ['记录ID', '文件名', '类型', '物体总数', '置信度阈值', '检测到的类别及数量', '备注', '状态', '检测时间']
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col, h in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    for row_idx, r in enumerate(records, 2):
        det_json = r.get('detection_json') or {}
        class_counts = det_json.get('class_counts', {})
        counts_str = '; '.join(f'{k}:{v}' for k, v in class_counts.items()) if class_counts else '无'

        values = [
            r['id'],
            r['file_name'],
            '图片' if r['file_type'] == 'image' else '视频',
            r['total_objects'],
            r['conf_threshold'],
            counts_str,
            r.get('remark') or '',
            r['status'],
            r['created_at'].strftime('%Y-%m-%d %H:%M:%S') if r.get('created_at') else '',
        ]
        for col, v in enumerate(values, 1):
            cell = ws1.cell(row=row_idx, column=col, value=v)
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center', wrap_text=True)

    # 列宽
    col_widths = [8, 30, 8, 10, 12, 40, 20, 10, 20]
    for i, w in enumerate(col_widths, 1):
        ws1.column_dimensions[chr(64 + i)].width = w

    # ---------- Sheet2: 类别统计 ----------
    ws2 = wb.create_sheet('类别数量统计')
    ws2.append(['类别名称', '出现次数（累计）'])
    for col in range(1, 3):
        cell = ws2.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border

    total_class_counts = {}
    for r in records:
        det_json = r.get('detection_json') or {}
        for k, v in det_json.get('class_counts', {}).items():
            total_class_counts[k] = total_class_counts.get(k, 0) + v

    sorted_counts = sorted(total_class_counts.items(), key=lambda x: x[1], reverse=True)
    for row_idx, (cls_name, cnt) in enumerate(sorted_counts, 2):
        ws2.cell(row=row_idx, column=1, value=cls_name).border = thin_border
        ws2.cell(row=row_idx, column=2, value=cnt).border = thin_border

    ws2.column_dimensions['A'].width = 20
    ws2.column_dimensions['B'].width = 18

    # 写入内存
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
